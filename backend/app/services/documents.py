from __future__ import annotations

import os
import re
import uuid
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import List

from bs4 import BeautifulSoup

from ..config import CHUNK_MAX_CHARS, UPLOAD_DIR
from ..db import connection_scope, utcnow
from .chunking import Chunk, Segment, build_segments_from_lines, chunk_segments, normalize_text

try:
    import fitz  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    fitz = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - optional dependency
    DocxDocument = None


HTML_HEADING_RE = re.compile(r"^h[1-6]$")


@dataclass
class ParsedChunk:
    id: str
    document_id: str
    chunk_index: int
    section_path: str
    page_no: int | None
    content: str
    created_at: str


def store_uploaded_file(file_bytes: bytes, filename: str) -> tuple[str, Path]:
    document_id = str(uuid.uuid4())
    safe_name = f"{document_id}_{Path(filename).name}"
    target = UPLOAD_DIR / safe_name
    target.write_bytes(file_bytes)
    return document_id, target


def create_document_record(document_id: str, filename: str, stored_path: Path) -> None:
    now = utcnow()
    extension = stored_path.suffix.lower()
    with connection_scope() as conn:
        conn.execute(
            """
            INSERT INTO documents (id, filename, stored_path, extension, status, chunk_count, uploaded_at, updated_at)
            VALUES (?, ?, ?, ?, ?, 0, ?, ?)
            """,
            (document_id, filename, str(stored_path), extension, "processing", now, now),
        )


def set_document_status(document_id: str, status: str, chunk_count: int = 0) -> None:
    with connection_scope() as conn:
        conn.execute(
            "UPDATE documents SET status = ?, chunk_count = ?, updated_at = ? WHERE id = ?",
            (status, chunk_count, utcnow(), document_id),
        )


def list_documents() -> list[dict]:
    with connection_scope() as conn:
        rows = conn.execute(
            "SELECT id, filename, status, chunk_count, uploaded_at, updated_at FROM documents ORDER BY uploaded_at DESC"
        ).fetchall()
        return [
            {
                "id": row["id"],
                "filename": row["filename"],
                "status": row["status"],
                "chunkCount": row["chunk_count"],
                "uploadedAt": row["uploaded_at"],
                "updatedAt": row["updated_at"],
            }
            for row in rows
        ]


def get_document(document_id: str) -> dict | None:
    with connection_scope() as conn:
        row = conn.execute("SELECT * FROM documents WHERE id = ?", (document_id,)).fetchone()
        return dict(row) if row else None


def delete_document_record(document_id: str) -> None:
    with connection_scope() as conn:
        conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))
        conn.execute("DELETE FROM documents WHERE id = ?", (document_id,))


def replace_document_chunks(document_id: str, chunks: list[ParsedChunk]) -> None:
    with connection_scope() as conn:
        conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))
        conn.executemany(
            """
            INSERT INTO chunks (id, document_id, chunk_index, section_path, page_no, content, created_at)
            VALUES (:id, :document_id, :chunk_index, :section_path, :page_no, :content, :created_at)
            """,
            [asdict(chunk) for chunk in chunks],
        )
        conn.execute(
            "UPDATE documents SET status = ?, chunk_count = ?, updated_at = ? WHERE id = ?",
            ("ready", len(chunks), utcnow(), document_id),
        )


def fetch_all_chunks() -> list[dict]:
    with connection_scope() as conn:
        rows = conn.execute(
            """
            SELECT c.id, c.document_id, c.chunk_index, c.section_path, c.page_no, c.content, d.filename
            FROM chunks c
            JOIN documents d ON d.id = c.document_id
            ORDER BY d.uploaded_at DESC, c.chunk_index ASC
            """
        ).fetchall()
        return [dict(row) for row in rows]


def fetch_chunks_by_ids(chunk_ids: list[str]) -> list[dict]:
    if not chunk_ids:
        return []
    placeholders = ",".join("?" for _ in chunk_ids)
    with connection_scope() as conn:
        rows = conn.execute(
            f"""
            SELECT c.id, c.document_id, c.chunk_index, c.section_path, c.page_no, c.content, d.filename
            FROM chunks c
            JOIN documents d ON d.id = c.document_id
            WHERE c.id IN ({placeholders})
            """,
            chunk_ids,
        ).fetchall()
    row_map = {row["id"]: dict(row) for row in rows}
    return [row_map[chunk_id] for chunk_id in chunk_ids if chunk_id in row_map]


def parse_document_to_chunks(document_id: str, stored_path: str) -> list[ParsedChunk]:
    path = Path(stored_path)
    extension = path.suffix.lower()
    if extension == ".pdf":
        segments = parse_pdf(path)
    elif extension == ".docx":
        segments = parse_docx(path)
    elif extension in {".txt", ".md"}:
        segments = parse_text_like(path)
    elif extension == ".html":
        segments = parse_html(path)
    else:
        raise ValueError(f"暂不支持的文件类型: {extension}")

    chunks = chunk_segments(segments)
    if not chunks:
        raise ValueError("未从文档中解析到有效文本")
    return [
        ParsedChunk(
            id=str(uuid.uuid4()),
            document_id=document_id,
            chunk_index=index,
            section_path=chunk.section_path,
            page_no=chunk.page_no,
            content=chunk.content[: CHUNK_MAX_CHARS * 2],
            created_at=utcnow(),
        )
        for index, chunk in enumerate(chunks)
    ]


def parse_pdf(path: Path) -> List[Segment]:
    if fitz is None:
        raise ValueError("PyMuPDF 未安装，无法解析 PDF")
    doc = fitz.open(path)
    segments: List[Segment] = []
    for page_index, page in enumerate(doc, start=1):
        text = normalize_text(page.get_text())
        if not text:
            continue
        paragraphs = re.split(r"(?<=。)\s+|\n+", text)
        for paragraph in paragraphs:
            paragraph = normalize_text(paragraph)
            if paragraph:
                segments.append(Segment(text=paragraph, page_no=page_index))
    return segments


def parse_docx(path: Path) -> List[Segment]:
    if DocxDocument is None:
        raise ValueError("python-docx 未安装，无法解析 DOCX")
    doc = DocxDocument(path)
    segments: List[Segment] = []
    headings: List[str] = []
    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if not text:
            continue
        style_name = getattr(paragraph.style, "name", "")
        if style_name.lower().startswith("heading") or re.match(r"^标题\s*\d*", style_name):
            headings = headings[:2]
            headings.append(text)
            continue
        segments.append(Segment(text=text, section_path=" / ".join(headings), page_no=1))
    return segments


def parse_text_like(path: Path) -> List[Segment]:
    raw_text = path.read_text(encoding="utf-8", errors="ignore")
    return build_segments_from_lines(raw_text.splitlines())


def parse_html(path: Path) -> List[Segment]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    headings: List[str] = []
    segments: List[Segment] = []
    for node in soup.find_all([re.compile(r"h[1-6]"), "p", "li"]):
        text = normalize_text(node.get_text(" ", strip=True))
        if not text:
            continue
        if HTML_HEADING_RE.match(node.name or ""):
            level = int((node.name or "h1")[1])
            headings = headings[: max(level - 1, 0)]
            headings.append(text)
            continue
        segments.append(Segment(text=text, section_path=" / ".join(headings), page_no=1))
    return segments


def remove_file(path_str: str) -> None:
    try:
        os.remove(path_str)
    except FileNotFoundError:
        return
